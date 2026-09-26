import os
from pathlib import Path

import openpyxl
import psycopg
from dotenv import load_dotenv

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
EXCEL = BASE_DIR / "unisphere_database.xlsx"

DATABASE_URL = (
    os.getenv("DATABASE_URL")
    or os.getenv("POSTGRES_URL")
    or os.getenv("POSTGRES_PRISMA_URL")
)

if not DATABASE_URL:
    raise SystemExit(
        "Set DATABASE_URL or POSTGRES_URL first."
    )

if not EXCEL.exists():
    raise SystemExit(f"Missing {EXCEL}")


SCHEMAS = {
    "users": """
        id BIGINT PRIMARY KEY,
        name TEXT,
        email TEXT,
        password TEXT,
        role TEXT,
        department TEXT,
        class_year TEXT,
        roll_no TEXT,
        employee_id TEXT,
        phone TEXT,
        status TEXT
    """,

    "attendance": """
        id BIGINT PRIMARY KEY,
        student_id BIGINT,
        subject TEXT,
        date TEXT,
        present BOOLEAN
    """,

    "faculty_attendance": """
        id BIGINT PRIMARY KEY,
        faculty_id BIGINT,
        date TEXT,
        present BOOLEAN
    """,

    "assignments": """
        id BIGINT PRIMARY KEY,
        title TEXT,
        subject TEXT,
        description TEXT,
        due_date TEXT,
        max_marks INTEGER,
        created_by BIGINT,
        created_at TEXT
    """,

    "submissions": """
        id BIGINT PRIMARY KEY,
        assignment_id BIGINT,
        student_id BIGINT,
        submitted_at TEXT,
        file_name TEXT,
        status TEXT,
        marks DOUBLE PRECISION,
        feedback TEXT
    """,

    "notices": """
        id BIGINT PRIMARY KEY,
        title TEXT,
        body TEXT,
        category TEXT,
        author TEXT,
        created_at TEXT
    """,

    "results": """
        id BIGINT PRIMARY KEY,
        student_id BIGINT,
        subject TEXT,
        internal DOUBLE PRECISION,
        external DOUBLE PRECISION,
        total DOUBLE PRECISION,
        grade TEXT,
        published BOOLEAN
    """,

    "fees": """
        id BIGINT PRIMARY KEY,
        student_id BIGINT,
        semester_fee DOUBLE PRECISION,
        paid DOUBLE PRECISION
    """,

    "placements": """
        id BIGINT PRIMARY KEY,
        company TEXT,
        role TEXT,
        package TEXT,
        eligibility TEXT,
        deadline TEXT,
        status TEXT
    """,

    "timetable": """
        id BIGINT PRIMARY KEY,
        day TEXT,
        start_time TEXT,
        end_time TEXT,
        subject TEXT,
        faculty TEXT,
        room TEXT,
        class_year TEXT
    """,
}


def pg_value(table, col, value):
    if value is None:
        return None

    if col in {"present", "published"}:
        if isinstance(value, bool):
            return value

        if isinstance(value, str):
            return value.strip().lower() in {
                "true",
                "1",
                "yes",
                "y",
                "present",
            }

        return bool(value)

    return value


print("========================================")
print(" UniSphere Excel → Neon Migration")
print("========================================")
print()
print("Excel:", EXCEL)
print("Connecting to Neon/PostgreSQL...")

wb = openpyxl.load_workbook(
    EXCEL,
    read_only=True,
    data_only=True
)


with psycopg.connect(DATABASE_URL) as conn:

    with conn.cursor() as cur:

        print()
        print("Creating tables if necessary...")

        # Create tables
        for table, schema in SCHEMAS.items():
            cur.execute(
                f"CREATE TABLE IF NOT EXISTS {table} ({schema})"
            )

        conn.commit()

        print("Tables ready.")
        print()

        # The Neon database is currently being populated
        # from the original Excel database.
        #
        # WARNING:
        # TRUNCATE removes existing data from these tables.
        # This is intentional for the initial migration.

        print("Clearing existing table data...")

        for table in SCHEMAS:
            cur.execute(f"TRUNCATE TABLE {table}")

        conn.commit()

        print("Existing data cleared.")
        print()

        # Import every Excel sheet
        for table in SCHEMAS:

            print(f"Importing: {table}")

            if table not in wb.sheetnames:
                raise SystemExit(
                    f"Excel sheet '{table}' was not found."
                )

            ws = wb[table]

            header_row = next(
                ws.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True
                )
            )

            headers = [
                str(column).strip()
                for column in header_row
            ]

            # Make sure the Excel file contains the ID column
            if "id" not in headers:
                raise SystemExit(
                    f"Sheet '{table}' does not contain an 'id' column."
                )

            quoted_columns = ",".join(
                f'"{header}"'
                for header in headers
            )

            placeholders = ",".join(
                ["%s"] * len(headers)
            )

            sql = (
                f"INSERT INTO {table} "
                f"({quoted_columns}) "
                f"VALUES ({placeholders})"
            )

            count = 0
            batch = []

            for raw_row in ws.iter_rows(
                min_row=2,
                values_only=True
            ):

                values = tuple(
                    pg_value(table, header, value)
                    for header, value
                    in zip(headers, raw_row)
                )

                batch.append(values)

                if len(batch) >= 1000:

                    cur.executemany(
                        sql,
                        batch
                    )

                    count += len(batch)
                    batch = []

            # Insert remaining rows
            if batch:

                cur.executemany(
                    sql,
                    batch
                )

                count += len(batch)

            print(f"  ✓ {count:,} rows imported")

        conn.commit()


wb.close()

print()
print("========================================")
print(" Migration complete!")
print("========================================")